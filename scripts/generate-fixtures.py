#!/usr/bin/env python3
"""Generate reproducible synthetic acceptance documents, no real client data."""
from pathlib import Path
import argparse

def main():
    parser=argparse.ArgumentParser();parser.add_argument('output',type=Path);args=parser.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    from PIL import Image,ImageDraw,ImageFont
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen import canvas
    texts=['Power supply voltage: 12 VDC.','Connect J3 to the STM32F407 control board.','Maximum torque is 4.5 Nm.','Use the M4x20 screws.','Communication uses CAN and RS-485.']
    image=Image.new('RGB',(1600,600),'white');draw=ImageDraw.Draw(image);font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',36)
    for i,line in enumerate(texts):draw.text((60,60+90*i),line,font=font,fill='black')
    annotation=args.output/'annotated.png';image.save(annotation)
    doc=Document();doc.add_heading('Technical validation',0)
    for line in texts:doc.add_paragraph(line)
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Parameter';table.cell(0,1).text='Value';table.cell(1,0).text='Voltage';table.cell(1,1).text='12 VDC';doc.add_picture(str(annotation));doc.save(args.output/'technical.docx')
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(5)).text='\n'.join(texts);p.save(args.output/'technical.pptx')
    c=canvas.Canvas(str(args.output/'native.pdf'));c.setFont('Helvetica-Bold',18);c.drawString(40,790,'Technical validation');c.setFont('Helvetica',12)
    for i,line in enumerate(texts):c.drawString(40,750-i*28,line)
    c.drawImage(str(annotation),40,280,width=500,height=188);c.save()
    c=canvas.Canvas(str(args.output/'scanned.pdf'));c.drawImage(str(annotation),30,450,width=550,height=206);c.save()

if __name__=='__main__':main()
