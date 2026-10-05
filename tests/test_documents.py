import zipfile
from io import BytesIO
import pytest
from app.documents import validate,extract,exports

def test_docx_tables_images_preserved(tmp_path):
    from docx import Document
    from PIL import Image
    from pypdf import PdfReader
    image=tmp_path/'test.png';Image.new('RGB',(100,60),'blue').save(image)
    source=tmp_path/'input.docx';doc=Document();doc.add_heading('Control board',1)
    doc.add_paragraph('Connect J3 to STM32F407.')
    doc.add_picture(str(image));doc.add_table(rows=1,cols=2).cell(0,0).text='Power supply'
    doc.tables[0].cell(0,1).text='12 VDC';doc.save(source)
    validate(source,'.docx');texts=extract(source,'.docx')
    assert texts==['Control board','Connect J3 to STM32F407.','Power supply','12 VDC']
    targets=['Carte de commande','Connectez J3 à STM32F407.','Alimentation','12 VDC']
    exports(source,'.docx',targets,tmp_path/'out')
    translated=Document(tmp_path/'out/translation.docx')
    assert len(translated.inline_shapes)==1
    assert translated.tables[0].cell(0,0).text=='Alimentation'
    assert extract(tmp_path/'out/translation.docx','.docx')==targets
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(tmp_path/'out/translation.docx') as b:
        assert a.read('word/media/image1.png')==b.read('word/media/image1.png')
    assert '12 VDC' in PdfReader(tmp_path/'out/translation.pdf').pages[0].extract_text()

def test_pptx_roundtrip(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Control board'
    f=tmp_path/'input.pptx';p.save(f);validate(f,'.pptx')
    exports(f,'.pptx',['Carte de commande'],tmp_path/'out')
    assert extract(tmp_path/'out/translation.pptx','.pptx')==['Carte de commande']

def test_fake_pdf_refused(tmp_path):
    f=tmp_path/'fake.pdf';f.write_text('not PDF')
    with pytest.raises(ValueError):validate(f,'.pdf')

def test_macro_refused(tmp_path):
    f=tmp_path/'fake.docx'
    with zipfile.ZipFile(f,'w') as z:z.writestr('word/document.xml','');z.writestr('word/vbaProject.bin','x')
    with pytest.raises(ValueError):validate(f,'.docx')

def test_zip_bomb_refused(tmp_path):
    f=tmp_path/'fake.docx'
    with zipfile.ZipFile(f,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml','x'*(2*1024*1024))
    with pytest.raises(ValueError):validate(f,'.docx')
