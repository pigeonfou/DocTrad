from sqlalchemy import select
from app.db import Session,User,Job
from app.auth import password_hash
from app.config import DATA

def signin(client,name='alice',admin=False):
    with Session.begin() as db:db.add(User(username=name,password=password_hash('very-strong-test-password'),admin=admin))
    client.get('/login')
    from itsdangerous import TimestampSigner
    import json,base64
    from app.config import SECRET
    def token():
        return json.loads(base64.b64decode(TimestampSigner(SECRET).unsign(client.cookies['session'])))['csrf']
    response=client.post('/login',data=dict(username=name,password='very-strong-test-password',token=token()))
    assert response.status_code==200
    return token()

def test_auth_required(client):
    assert client.get('/').url.path=='/login'

def test_csrf_required(client):
    client.get('/login')
    assert client.post('/login',data={'username':'a','password':'b','token':'invalid'}).status_code==403

def test_home_and_admin_permissions(client):
    signin(client)
    assert 'Nouvelle traduction' in client.get('/').text
    assert client.get('/admin').status_code==403

def test_no_model_explicit(client):
    token=signin(client)
    response=client.post('/jobs',data={'token':token},files={'document':('a.pdf',b'%PDF-')})
    assert response.status_code==400
    assert 'Aucun modèle local' in response.text

def test_cross_user_job_hidden(client):
    signin(client)
    with Session.begin() as db:
        other=User(username='bob',password=password_hash('another-strong-password'));db.add(other);db.flush()
        db.add(Job(id='testjob',owner_id=other.id,filename='confidential',suffix='.pdf',model_id='opus',model_version='a'*40))
    assert client.get('/jobs/testjob').status_code==404
    assert client.get('/jobs/testjob/download/pdf').status_code==404

def test_xss_is_escaped(client):
    signin(client)
    with Session.begin() as db:
        user=db.scalar(select(User).where(User.username=='alice'))
        db.add(Job(id='xss',owner_id=user.id,filename='<script>alert(1)</script>',suffix='.pdf',model_id='opus',model_version='a'*40))
    body=client.get('/').text
    assert '&lt;script&gt;' in body
    assert '<script>alert(1)' not in body

def test_health(client):
    assert client.get('/health').status_code==200
    assert client.get('/health').json()['offline']

def test_runtime_secret_required(client):
    from fastapi.testclient import TestClient
    from app.runtime import app
    with TestClient(app) as c:
        assert c.post('/translate',json={'model_id':'opus','revision':'a'*40,'texts':['Hello']}).status_code==403

def test_upload_process_correct_download(client,tmp_path,monkeypatch):
    """Complete HTTP workflow with validated translation memory, no AI claim."""
    from app.registry import Registry
    from app.db import Memory,Segment,Correction
    from app.worker import claim,run_job
    from docx import Document
    import io
    selected={'id':'memory-test','revision':'a'*40}
    monkeypatch.setattr(Registry,'choose',lambda self,requested='auto':selected)
    token=signin(client)
    with Session.begin() as db:
        user=db.scalar(select(User).where(User.username=='alice'))
        db.add(Memory(owner_id=user.id,source='Power supply voltage: 12 VDC.',target='Tension alimentation : 12 VDC.'))
    document=Document();document.add_paragraph('Power supply voltage: 12 VDC.');stream=io.BytesIO();document.save(stream)
    result=client.post('/jobs',data={'token':token},files={'document':('../../danger.docx',stream.getvalue())})
    assert result.status_code==200
    job_id=result.url.path.split('/')[-1]
    assert claim()==job_id;run_job(job_id)
    assert client.get(f'/jobs/{job_id}/status').json()['status']=='done'
    assert client.get(f'/jobs/{job_id}/download/pdf').content.startswith(b'%PDF-')
    with Session() as db:seg=db.scalar(select(Segment).where(Segment.job_id==job_id));segment_id=seg.id
    result=client.post(f'/jobs/{job_id}/segments/{segment_id}',data={'token':token,'target':'Tension nominale : 12 VDC.'})
    assert result.status_code==200
    assert claim()==job_id;run_job(job_id)
    output=client.get(f'/jobs/{job_id}/download/docx')
    assert Document(io.BytesIO(output.content)).paragraphs[0].text=='Tension nominale : 12 VDC.'
    with Session() as db:assert db.scalar(select(Correction)).previous=='Tension alimentation : 12 VDC.'
    assert not (DATA/'danger.docx').exists()
