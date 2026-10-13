from app.config import DATA
from app.db import Session,User,Job,Segment,Memory
from app.worker import claim,run_job
from sqlalchemy import select

def test_queue_to_export_using_validated_memory():
    from docx import Document
    DATA.joinpath('uploads').mkdir(parents=True,exist_ok=True)
    source='Power supply voltage: 12 VDC.';target='Tension alimentation : 12 VDC.'
    with Session.begin() as db:
        user=User(username='alice',password='not-used');db.add(user);db.flush()
        db.add(Job(id='memory-job',owner_id=user.id,filename='a.docx',suffix='.docx',model_id='opus',model_version='a'*40))
        db.add(Memory(owner_id=user.id,source=source,target=target))
    doc=Document();doc.add_paragraph(source);doc.save(DATA/'uploads/memory-job.docx')
    assert claim()=='memory-job'
    assert claim() is None
    run_job('memory-job')
    with Session() as db:
        assert db.get(Job,'memory-job').status=='done'
        assert db.scalar(select(Segment)).target==target
    assert (DATA/'outputs/memory-job/translation.pdf').is_file()

def test_export_after_correction():
    test_queue_to_export_using_validated_memory()
    with Session.begin() as db:
        db.get(Job,'memory-job').status='export_queued'
        db.scalar(select(Segment)).target='Alimentation nominale : 12 VDC.'
    assert claim()=='memory-job';run_job('memory-job')
    from docx import Document
    assert Document(DATA/'outputs/memory-job/translation.docx').paragraphs[0].text=='Alimentation nominale : 12 VDC.'


def test_known_layout_failure_recovered_without_losing_correction():
    from reportlab.pdfgen import canvas
    from app.worker import recover_layout_failures
    DATA.joinpath('uploads').mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(DATA/'uploads/layout-recovery.pdf'))
    c.drawString(60,700,'Agency');c.save()
    with Session.begin() as db:
        user=User(username='layout-review',password='unused');db.add(user);db.flush()
        db.add(Job(id='layout-recovery',owner_id=user.id,filename='reference.pdf',suffix='.pdf',
                   model_id='opus',model_version='a'*40,status='failed',error='Coordonnées OCR hors de la page'))
        db.add(Segment(job_id='layout-recovery',position=0,source='Agency',machine='Agence',
                       target='Agence relue',warnings=[],metadata_json={'page':1,'box':[114,1942,113,1963]}))
    recover_layout_failures()
    assert claim() == 'layout-recovery'
    run_job('layout-recovery')
    with Session() as db:
        job = db.get(Job,'layout-recovery')
        segment = db.scalar(select(Segment).where(Segment.job_id == job.id))
        assert job.status == 'done' and not job.error
        assert segment.target == 'Agence relue'
        assert any('Coordonnées OCR invalides' in warning for warning in segment.warnings)
    recover_layout_failures()
    assert claim() is None
    assert (DATA/'outputs/layout-recovery/translation.pdf').is_file()


def test_postgres_claim_is_atomic():
    import pytest
    from app.db import engine
    if engine.dialect.name != "postgresql":pytest.skip("Native PostgreSQL CI test")
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    with Session.begin() as db:
        user=User(username="queue-test",password="unused");db.add(user);db.flush()
        for i in range(4):db.add(Job(id=f"concurrent-{i}",owner_id=user.id,filename="test.docx",suffix=".docx",model_id="opus",model_version="a"*40))
    gate=Barrier(4)
    def reserve():gate.wait();return claim()
    with ThreadPoolExecutor(max_workers=4) as pool:results=list(pool.map(lambda _:reserve(),range(4)))
    assert len(set(results))==4 and None not in results
    assert claim() is None



def test_worker_sends_context_and_audits_changed_units(monkeypatch):
    from docx import Document
    from app import worker
    source = 'Shadow Radio operates over 26 miles.'
    target = 'La radio fantôme fonctionne sur 26 kilomètres.'
    DATA.joinpath('uploads').mkdir(parents=True, exist_ok=True)
    with Session.begin() as db:
        user = User(username='context-worker', password='unused'); db.add(user); db.flush()
        db.add(Job(id='context-job', owner_id=user.id, filename='context.docx', suffix='.docx', model_id='opus', model_version='a'*40))
    doc = Document(); doc.add_paragraph(source); doc.save(DATA/'uploads/context-job.docx')
    class Client:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def post(self, url, *, headers, json):
            assert json['texts'] == [source]
            assert json['contextual'] is True
            assert json['glossary']['Shadow Radio'] == 'Shadow Radio'
            return type('Response', (), {'status_code':200, 'json':lambda self: {'texts':[target]}})()
    monkeypatch.setattr(worker.httpx, 'Client', Client)
    assert claim() == 'context-job'
    run_job('context-job')
    with Session() as db:
        assert db.get(Job, 'context-job').status == 'done'
        segment = db.scalar(select(Segment).where(Segment.job_id == 'context-job'))
        assert segment.target == target
        assert any('26 miles' in warning for warning in segment.warnings)
        assert any('Shadow Radio' in warning for warning in segment.warnings)



def test_remote_job_waits_when_another_worker_holds_bc250(monkeypatch):
    import fcntl
    from app import worker
    DATA.mkdir(parents=True,exist_ok=True)
    with Session.begin() as db:
        user=User(username='remote-lock',password='unused');db.add(user);db.flush()
        db.add(Job(id='remote-lock-job',owner_id=user.id,filename='a.pdf',suffix='.pdf',
                   model_id='remote-bc250',model_version='a'*40,status='running'))
    calls=[]
    monkeypatch.setattr(worker,'_run_job',lambda job_id:calls.append(job_id))
    monkeypatch.setattr(worker.time,'sleep',lambda seconds:None)
    with open(DATA/'remote-job.lock','a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX | fcntl.LOCK_NB)
        run_job('remote-lock-job')
    assert not calls
    with Session() as db:assert db.get(Job,'remote-lock-job').status=='queued'
    run_job('remote-lock-job')
    assert calls==['remote-lock-job']


def test_defective_memories_are_retranslated_without_deleting_user_entries():
    from app.worker import reusable_translation
    from app.terminology import terminology
    assert reusable_translation('Windows CPA application', 'Application Windows CPA', terminology())
    assert not reusable_translation('Compact dual-channel device', 'Appareil compact 双通道', terminology())
    assert not reusable_translation('Windows CPA application', 'Application Windows CPA\n["Windows", "CPA"]', terminology())
    assert not reusable_translation('Internal SAASM GPS', 'GPS interne', terminology())
